"""Build the Chinese method survey workbook from the translated triage / literature JSON and the measured probes."""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

S = Path('/private/tmp/claude-501/-Users-slian-Desktop-3D-Worldedge-OptiScene/5fdb3dba-c432-499d-ad8b-3e4a495041ec/scratchpad')
OUT = Path(sys.argv[1])

# ---- sheet 1: ranked methods (written by hand from probes, literature and the 23 refuter reviews) ----
RANK_HEAD = ['排序', '分组', '方法', '借鉴自', '改什么', '针对的失败', '证据（文献数字或我们实测）', '预期收益', '成本', '状态']
RANK = [
    [1, 'A 已实测有效，CPU，马上能用', 'S2 确定性重摆 + 关系保护',
     'PolyLayout（表 56）；同类：Holodeck DFS（41）、Roomer（59）',
     'S1 之后加一级：好物体当障碍，坏物体从大到小沿墙、换墙重新锚定；尺寸、朝向、父物体不变；靠墙物体只在仍靠墙的位置里选',
     'S1 后仍有碰撞/越界的房（test 2,225 间）',
     '实测 dev 600 间：checker 0.820 → 0.963，GT 关系保持 0.746 → 0.750；test 上 300 间 S1 失败房 0 → 0.64',
     'checker +0.14（dev）；可替掉大部分 checker 触发的 O2 调用', '0 GPU；约 0.9 s/房', '已实测（dev），待留出集确认'],
    [2, 'A 已实测有效，CPU，马上能用', 'S2 把 A4 同父重叠纳入修复目标',
     'LabBuilder（11）的“违规归因→修复”思路；提案 H1/R2（3 人质疑均未驳回）',
     '对 strict 失败的房再跑一遍 S2，目标里加入 A4；checker 和 strict 都不变差才采用',
     'S1 后过 checker 但 strict 失败（A4 占 69.5%）',
     '实测 dev 600 间：strict 0.355 → 0.572，checker 0.970，关系保持 0.753',
     'strict +0.22（dev）', '0 GPU；约 0.5 s/房',
     '已实测（dev）；代价：平均移动 31% 物体，位移中位 0.94 m，会把“椅子塞在桌下”推开，需要渲染抽查'],
    [3, 'A 已实测有效，CPU，马上能用', '门口可达修复（A3）',
     'LabBuilder（11）；RoboLayout（49）',
     '在可通行栅格上从门做可达搜索，把挡路的物体挪开',
     'strict 里的 A3 门不可达（dev S2 后 177 间含 A3）',
     '实测 test：只剩 A3 的 307 间修好 46%，全 test strict +0.016，上限 +0.036，无房间变差',
     'strict +0.02~0.04', '0 GPU', '已实测（test 子集）；待接到 S2 后在 dev 上跑'],
    [4, 'B 证据强，需 GPU 小实验，不训练', '可行性掩码解码',
     'Global-Local tree search（文献 L2）；检查器在环逐件放置；提案 P1（3 人质疑均未驳回）',
     '推理时用 checker 的规则屏蔽会碰撞/越界的坐标 token，模型只能输出可行位置',
     '原始输出碰撞（失败房 86.5% 有主碰撞）、抄上一条',
     '质疑后估计：原始 checker 0.443 → 0.55~0.68；S1 后 0.742 → 0.77~0.83',
     '让模型本身变好，不靠后处理；S2 之后的边际收益要重新测', '试点 1~2 GPU-h，不训练', '未做'],
    [5, 'B 证据强，需 GPU 小实验，不训练', '门禁区作为固定物体输入',
     'PolyLayout（56）门窗禁区；提案 H2/R4',
     '用 downstream.door_strips 生成门前禁区，作为 fixed 放进请求',
     '有门房 strict 0.225 vs 无门房 0.498',
     '未经质疑估计：strict +0.025~0.06', 'strict +0.03~0.06', '约 1 GPU-h 验证', '未做'],
    [6, 'B 证据强，需 GPU 或 API', '大房间分块生成，已放物体作为固定上下文',
     'HDSL（14）、STABLE（18）、FuncRoom-Agent（7）；提案 P5',
     '>20 件的房分块输出：每块先过 S1，再作为 fixed 进入下一块',
     '21+ 件房原始 checker 0.024，复制概率随输出位置升到 0.46~0.63',
     'FuncRoom-Agent 64 件/房碰撞约 10.7%；我们实测：前 8 条里已有一半复制，分块只能去掉约一半；分阶段 S1 在 110 间上 0.273 → 0.282',
     '估计 21~50 件 checker 0.024 → 0.15~0.35（未复核）', '3~4 天工程 + 1 GPU-h 试点', '部分探针已做；收益不确定'],
    [7, 'B 证据强，需 GPU 或 API', '只重生成失败物体（局部重生成）',
     'HDSL（14）固定上下文；Roomer（59）单违规编辑；提案 R4',
     '合法物体锁定为 fixed，只让模型重摆失败的几件',
     'S1 后仍失败的 11~50 件房',
     'Roomer：单违规编辑修好 98.96% vs 整房重写 32.48%；我们的 100 次 gpt-6.1-sol 测试已备好（对照：免费求解器 checker 0.74）',
     '若不明显超过 0.74 就不值得训练', '100 次 API 调用', '待你 clifton auth'],
    [8, 'C 需要训练（中期）', '坐标改成整数网格 / 子物体用父物体相对坐标',
     'SpatialGrammar（47）、NaLA、R3L（文献 L4）；提案 P4',
     '输出格式改成网格整数，子物体写成在父物体坐标系里的偏移，再在 contract 里转回绝对坐标',
     '碰撞、子物体抄父物体位置、大房间序列太长',
     'SpatialGrammar：原始 JSON 碰撞 66.7% → 网格 13.6%；NaLA 越界 9.25 → 1.16；质疑后估计我们原始 +0.02~0.05，S1 后 +0.01~0.03',
     '+0.01~0.05', '2 天工程 + 一次 8B 重训', '未做'],
    [9, 'C 需要训练（中期）', '可验证奖励强化学习（GRPO，分项归一化）',
     'LayoutDSL（54）、ZoneMaestro（8）、SPG-Layout、Architect-Ant（文献 L1/L5）；提案 TS-4',
     '用 checker/strict 分项奖励训练 27B，按 GDPO 做分项归一化，并加尺寸/数量保护',
     '原始碰撞、门可达、非矩形房越界',
     'ZoneMaestro 碰撞 0.18 → 0.04；LayoutDSL 可达 0.58 → 0.96；质疑后估计我们 S1 后 +0.01~0.04',
     '+0.01~0.04（S1 后）', '每轮 5~6 GPU-h 采样 + 训练', '未做；收益偏小'],
    [10, 'C 需要训练（中期）', '同底座 thinking → 非 thinking 自蒸馏',
     'ThinkOPD（arXiv 2609.37044）；提案 TS-6',
     '同一个 27B 开 thinking 当老师，在线蒸馏给不开 thinking 的 LoRA',
     'thinking 0.79~0.92 vs 微调不 thinking 0.443',
     '质疑后估计：原始 +0.03~0.08（几乎都在 ≤20 件），S1 后 +0~0.02',
     'S1 后 +0~0.02', '80~130 GPU-h', '未做；不优先'],
    [11, 'D 需要你决定', '大房间允许受限缩小尺寸',
     '提案 H4；文献 L1：只用几何奖励会让尺寸缩小，需要尺寸保真项',
     '摆不下时，违规物体在先验 25 分位内缩小，再重排',
     '先验尺寸下 GT 摆法本身就放不下：21~50 件 0.770 → 0.124，51+ 0.790 → 0.061',
     '未经质疑估计：21~50 件 S1 后 0.402 → 0.45~0.55', '+0.05~0.15（21~50 件）', '0~3 GPU-h', '要改尺寸，需要你同意'],
    [12, 'E 不建议', 'DPO 再改一版（加关系过滤）', '—', '—', '—', '两位质疑者驳回：S1 不破坏关系，关系下降来自训练本身', '—', '—', '不做'],
    [13, 'E 不建议', '训练 27B 看违规报告再改（第二轮）', 'CYCLE（代码领域）', '—', '—', '两位质疑者驳回：可用训练目标太少，修复率估计只有 5~12%', '—', '—', '不做'],
    [14, 'E 不建议', '只加实例编号 / 解码禁重复', 'HSM、I-Scene', '—', '复制', '实测：删光全部复制，上限也只有原始 +0.012 / S1 后 +0.039', '—', '—', '只作附带项'],
    [15, 'E 不建议', '扩散/流匹配布局头', '文献 L3（26 篇）', '—', '—', '2025~26 年没有一篇在 >20 件、带门多边形房上报告物理通过率', '—', '—', '不做'],
    [16, 'E 不建议', '掩码重填（随机候选重放被掩码物体）', 'Steerable Scene Gen（36）', '—', '—', '实测 checker 0.71，但平均只放下 88% 物体，违反“一件不能丢”', '—', '—', '不做'],
]

# ---- sheet: the user's curated 10 (Downloads/Task_to_Room_文献与平台对比.xlsx) mapped to our verdicts ----
CURATED = [
    ['01', 'EmbodiedSmith', '02'], ['02', 'EmbodiedGen V2', '01'], ['03', 'FuncRoom-Agent', '07'],
    ['04', 'SceneMosaic', None], ['05', 'WEDGE', '26'], ['06', 'CoDA', '27'], ['07', 'SceneSmith', '04'],
    ['08', 'SAGE', '03'], ['09', 'SimWorld Studio', '30'], ['10', 'SceneNAT', '53'],
]
SCENEMOSAIC = {
    'verdict_zh': '思路可借鉴', 'method_zh': '参考图 → 逐物体资产与布局初稿 → 按支撑/容纳关系分成局部单元 → 物理稳定化和 VLM 局部修复；局部单元的变体组合成多样房间',
    'usable_zh': '“按局部单元独立修复、再组合”可用于 S2/局部重生成：每个支撑子树单独出候选，再挑互不冲突的组合（与提案 TS-2 子树重组同类，质疑后部署收益 +0.03~0.07）',
    'targets_zh': 'S1 后的多违规房', 'code_zh': '代码公开，依赖外部权重与 VLM 服务',
    'why_zh': '初稿来自图像和 SAM3D，不能直接替换 FastFill；只依据摘要和你的表，未读全文',
}

MEASURED = {  # triage index -> our measured probe result
    '56': 'S2：dev 600 间 checker 0.820 → 0.963；加 A4 后 strict 0.355 → 0.572',
    '11': '门可达修复：test strict +0.016（上限 +0.036）',
    '59': '移动阶梯：S1 失败房修好 33%（S1 自身候选 0.8%）',
    '41': 'DFS 去重放置：24 间试跑 checker 0.125 → 0.167，S1 单独 0.50；不如 S2',
    '14': '复制风险随位置上升；分块只能去掉约一半复制；100 次 API 局部重生成待跑',
    '18': '分阶段 S1：110 间 0.273 → 0.282，无明显收益',
    '36': '掩码重填：checker 0.71 但丢 12% 物体',
    '48': '可微优化修复：8 间试跑 0.625，与 S1 持平',
    '54': 'GT 转写到 LayoutDSL 锚点表示：8,647 间无解码失败，漂移 <2 cm',
}


def style(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDE6F0')
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'


def main():
    triage = json.load(open(S / 'triage_zh_01_35.json')) + json.load(open(S / 'triage_zh_36_70.json'))
    lit = json.load(open(S / 'lit_zh_000_064.json')) + json.load(open(S / 'lit_zh_065_129.json'))
    assert len(triage) == 70 and len(lit) == 130, (len(triage), len(lit))
    by = {t['index']: t for t in triage}
    wb = Workbook()
    ws = wb.active
    ws.title = '能提升的方法（排序）'
    ws.append(RANK_HEAD)
    for r in RANK:
        ws.append(r)
    style(ws, [5, 16, 22, 26, 36, 24, 40, 18, 16, 20])

    ws = wb.create_sheet('你的70条')
    ws.append(['编号', '名称', '结论', '优先级(0-7)', '方法', '可拿来的部分', '针对我们哪个失败', '我们实测', '符合接口约束', '代码', '理由', '链接'])
    for t in sorted(triage, key=lambda t: (-int(t['priority'] or 0), t['index'])):
        ws.append([t['index'], t['name'], t['verdict_zh'], t['priority'], t['method_zh'], t['usable_zh'], t['targets_zh'],
                   MEASURED.get(t['index'], ''), '是' if t['fits_contract'] else '否', t['code_zh'], t['why_zh'], t['url']])
    style(ws, [5, 26, 10, 8, 36, 40, 24, 30, 8, 18, 40, 30])

    ws = wb.create_sheet('你的精选10条')
    ws.append(['精选编号', '名称', '对应70条编号', '结论', '可拿来的部分', '针对我们哪个失败', '理由'])
    for k, name, idx in CURATED:
        t = by[idx] if idx else SCENEMOSAIC
        ws.append([k, name, idx or '不在70条里（本轮补）', t['verdict_zh'], t['usable_zh'], t['targets_zh'], t['why_zh']])
    style(ws, [8, 18, 12, 10, 50, 26, 50])

    ws = wb.create_sheet('补充文献130篇')
    ws.append(['角度', '标题', '年份', '方法', '报告结果', '对 FastFill 的用处', '已核对', '链接'])
    for p in sorted(lit, key=lambda p: (p['angle'], p['pos'])):
        ws.append([p['angle_zh'], p['title'], p['year'], p['method_zh'], p['result_zh'], p['use_zh'],
                   '是' if p['verified'] in (True, 'true', 'yes') else str(p['verified']), p['url']])
    style(ws, [16, 34, 8, 40, 30, 44, 8, 30])

    ws = wb.create_sheet('说明')
    for line in [
        ['来源：2026-10-20/21 两个调研 workflow 的原始记录（probes-20261021/workflow-extracts/）+ 本地探针实测 + 本表整理（2026-10-21）。'],
        ['“实测”都在 27B 已有输出上，只用 CPU；dev = dev-select 600 间（选模用过），test = 8,647 间（多次使用过），结论要在没碰过的房屋组留出集上确认。'],
        ['“质疑后估计”来自 23 份独立质疑意见；90 份计划只完成 23 份，没经过质疑的标“未经质疑”或“未复核”。'],
        ['checker = 碰撞/越界/支撑几何检查；strict = 下游准入（A2 门禁区、A3 门可达、A4 同父重叠）；S1 = 现有 CPU 平移修复；O2 = 升级到大模型。'],
        ['详细报告：docs/research/20261013-unknown-size-comparison/method-improvement-research-20261021.md'],
    ]:
        ws.append(line)
    ws.column_dimensions['A'].width = 140
    wb.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
